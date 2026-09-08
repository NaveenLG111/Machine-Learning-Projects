
import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.model_selection import train_test_split, KFold, cross_val_score
from sklearn.compose import ColumnTransformer
from sklearn.pipeline import Pipeline
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.linear_model import Ridge
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import ScatterChart, BarChart, Reference, Series


class ExcelWriterML:
    def __init__(self, file_path, output_excel='ml_model_report.xlsx', output_dir='plots', target=None, test_size=0.2, random_state=42):
        self.file_path = file_path
        self.output_excel = output_excel
        self.output_dir = output_dir
        self.target = target
        self.test_size = test_size
        self.random_state = random_state
        self.df = None
        self.X = None
        self.y = None
        self.X_train = None
        self.X_test = None
        self.y_train = None
        self.y_test = None
        self.model = None
        self.pipeline = None
        self.train_pred = None
        self.test_pred = None
        self.cv_r2 = None
        self.summary_df = None
        self.eda_missing_df = None
        self.eda_stats_df = None
        self.actual_pred_df = None
        self.output_metrics_df = None
        self.feature_importance_df = None
        self.numeric_df = None
        self.plot_files = []
        os.makedirs(self.output_dir, exist_ok=True)
        self._init_styles()

    def _init_styles(self):
        self.hdr_fill = PatternFill('solid', fgColor='1F4E78')
        self.hdr_font = Font(name='Calibri', bold=True, color='FFFFFF')
        self.body_font = Font(name='Calibri', size=11)
        self.title_font = Font(name='Calibri', size=14, bold=True)
        self.center = Alignment(horizontal='center', vertical='center')
        self.left = Alignment(horizontal='left', vertical='center')
        self.right = Alignment(horizontal='right', vertical='center')
        thin = Side(style='thin', color='D9D9D9')
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def load_data(self):
        self.df = pd.read_csv(self.file_path)
        if self.target is None:
            priority = ['target', 'label', 'y', 'price', 'cost', 'amount', 'sales', 'yield', 'production']
            for p in priority:
                matches = [c for c in self.df.columns if c.lower() == p or c.lower().endswith('_' + p) or p in c.lower()]
                if matches:
                    numeric_matches = [c for c in matches if pd.api.types.is_numeric_dtype(self.df[c])]
                    self.target = numeric_matches[0] if numeric_matches else matches[0]
                    break
            if self.target is None:
                numeric_candidates = [c for c in self.df.columns if pd.api.types.is_numeric_dtype(self.df[c]) and self.df[c].nunique(dropna=True) > 10]
                self.target = numeric_candidates[-1] if numeric_candidates else self.df.columns[-1]
        self.X = self.df.drop(columns=[self.target])
        self.y = self.df[self.target]
        self.numeric_df = self.df.select_dtypes(include=np.number)

    def build_model(self):
        num_cols = [c for c in self.X.columns if pd.api.types.is_numeric_dtype(self.X[c])]
        cat_cols = [c for c in self.X.columns if c not in num_cols]
        pre = ColumnTransformer([
            ('num', Pipeline([('imputer', SimpleImputer(strategy='median')), ('scaler', StandardScaler())]), num_cols),
            ('cat', Pipeline([('imputer', SimpleImputer(strategy='most_frequent')), ('onehot', OneHotEncoder(handle_unknown='ignore'))]), cat_cols)
        ])
        self.model = Ridge(alpha=10.0)
        self.pipeline = Pipeline([('pre', pre), ('model', self.model)])
        self.X_train, self.X_test, self.y_train, self.y_test = train_test_split(
            self.X, self.y, test_size=self.test_size, random_state=self.random_state
        )
        self.pipeline.fit(self.X_train, self.y_train)
        self.train_pred = self.pipeline.predict(self.X_train)
        self.test_pred = self.pipeline.predict(self.X_test)
        cv = KFold(n_splits=5, shuffle=True, random_state=self.random_state)
        self.cv_r2 = cross_val_score(self.pipeline, self.X, self.y, cv=cv, scoring='r2').mean()

        feature_names = self.pipeline.named_steps['pre'].get_feature_names_out()
        coef = np.abs(np.ravel(self.pipeline.named_steps['model'].coef_))
        self.feature_importance_df = pd.DataFrame({'Feature': feature_names, 'Importance': coef}).sort_values('Importance', ascending=False).head(20)

    def prepare_sheets_data(self):
        self.summary_df = pd.DataFrame({
            'Item': ['Target Column', 'Rows', 'Columns', 'Train Rows', 'Test Rows', 'Best Model', 'CV Mean R2'],
            'Value': [self.target, len(self.df), len(self.df.columns), len(self.X_train), len(self.X_test), 'Ridge(alpha=10.0)', self.cv_r2]
        })
        self.eda_missing_df = pd.DataFrame({
            'Column': self.df.columns,
            'Missing_Count': self.df.isna().sum().values,
            'Missing_%': (self.df.isna().mean() * 100).round(2).values,
            'Dtype': self.df.dtypes.astype(str).values
        })
        self.eda_stats_df = self.df.describe(include='all').transpose().reset_index().rename(columns={'index': 'Column'})
        self.actual_pred_df = self.X_test.reset_index(drop=True).copy()
        self.actual_pred_df['Actual'] = self.y_test.reset_index(drop=True)
        self.actual_pred_df['Predicted'] = self.test_pred
        self.actual_pred_df['Residual'] = self.actual_pred_df['Actual'] - self.actual_pred_df['Predicted']
        self.actual_pred_df['Abs_Error'] = self.actual_pred_df['Residual'].abs()
        self.actual_pred_df['APE_%'] = np.where(
            self.actual_pred_df['Actual'] != 0,
            self.actual_pred_df['Abs_Error'] / self.actual_pred_df['Actual'].abs() * 100,
            np.nan
        )
        self.output_metrics_df = pd.DataFrame([
            ['Train', 'Ridge(alpha=10.0)', mean_absolute_error(self.y_train, self.train_pred), mean_squared_error(self.y_train, self.train_pred) ** 0.5, r2_score(self.y_train, self.train_pred)],
            ['Test', 'Ridge(alpha=10.0)', mean_absolute_error(self.y_test, self.test_pred), mean_squared_error(self.y_test, self.test_pred) ** 0.5, r2_score(self.y_test, self.test_pred)],
            ['CV Mean', 'Ridge(alpha=10.0)', np.nan, np.nan, self.cv_r2]
        ], columns=['Split', 'Model', 'MAE', 'RMSE', 'R2'])

    def save_all_plots(self):
        sns.set_theme(style='whitegrid')

        def save(fig, name):
            path = os.path.join(self.output_dir, name)
            fig.tight_layout()
            fig.savefig(path, dpi=200, bbox_inches='tight')
            plt.close(fig)
            self.plot_files.append(path)

        fig, ax = plt.subplots(figsize=(10, 6))
        sns.histplot(self.df[self.target], kde=True, ax=ax)
        ax.set_title(f'Target Distribution - {self.target}')
        save(fig, 'target_distribution.png')

        if len(self.numeric_df.columns) >= 2:
            fig, ax = plt.subplots(figsize=(10, 8))
            corr = self.numeric_df.corr(numeric_only=True)
            sns.heatmap(corr, cmap='coolwarm', annot=False, ax=ax)
            ax.set_title('Correlation Heatmap')
            save(fig, 'correlation_heatmap.png')

        num_features = [c for c in self.X.columns if pd.api.types.is_numeric_dtype(self.X[c])]
        if len(num_features) > 0:
            feat = num_features[0]
            fig, ax = plt.subplots(figsize=(10, 6))
            sns.scatterplot(x=self.df[feat], y=self.df[self.target], ax=ax)
            ax.set_title(f'{feat} vs {self.target}')
            save(fig, 'feature_vs_target.png')

        fig, ax = plt.subplots(figsize=(8, 8))
        sns.scatterplot(x=self.actual_pred_df['Actual'], y=self.actual_pred_df['Predicted'], ax=ax)
        mn = min(self.actual_pred_df['Actual'].min(), self.actual_pred_df['Predicted'].min())
        mx = max(self.actual_pred_df['Actual'].max(), self.actual_pred_df['Predicted'].max())
        ax.plot([mn, mx], [mn, mx], 'r--')
        ax.set_title('Actual vs Predicted')
        save(fig, 'actual_vs_predicted.png')

        fig, ax = plt.subplots(figsize=(10, 6))
        sns.histplot(self.actual_pred_df['Residual'], kde=True, ax=ax)
        ax.axvline(0, color='red', linestyle='--')
        ax.set_title('Residual Distribution')
        save(fig, 'residual_distribution.png')

        fig, ax = plt.subplots(figsize=(10, 6))
        sns.scatterplot(x=self.actual_pred_df['Predicted'], y=self.actual_pred_df['Residual'], ax=ax)
        ax.axhline(0, color='red', linestyle='--')
        ax.set_title('Residuals vs Predicted')
        save(fig, 'residuals_vs_predicted.png')

        fig, ax = plt.subplots(figsize=(8, 5))
        sns.barplot(data=self.output_metrics_df[self.output_metrics_df['Split'] != 'CV Mean'], x='Split', y='R2', ax=ax)
        ax.set_title('Train vs Test R2')
        save(fig, 'train_test_r2.png')

        fig, ax = plt.subplots(figsize=(10, 8))
        sns.barplot(data=self.feature_importance_df.head(15), y='Feature', x='Importance', ax=ax)
        ax.set_title('Top Feature Importance')
        save(fig, 'feature_importance.png')

        fig, axes = plt.subplots(2, 2, figsize=(14, 10))
        sns.histplot(self.df[self.target], kde=True, ax=axes[0, 0])
        axes[0, 0].set_title('Target Distribution')
        sns.scatterplot(x=self.actual_pred_df['Actual'], y=self.actual_pred_df['Predicted'], ax=axes[0, 1])
        axes[0, 1].plot([mn, mx], [mn, mx], 'r--')
        axes[0, 1].set_title('Actual vs Predicted')
        sns.histplot(self.actual_pred_df['Residual'], kde=True, ax=axes[1, 0])
        axes[1, 0].axvline(0, color='red', linestyle='--')
        axes[1, 0].set_title('Residual Distribution')
        sns.barplot(data=self.output_metrics_df[self.output_metrics_df['Split'] != 'CV Mean'], x='Split', y='R2', ax=axes[1, 1])
        axes[1, 1].set_title('Train vs Test R2')
        save(fig, 'model_summary_dashboard.png')

    def _write_df(self, ws, df, start_row=2, start_col=2, table_name='Table1'):
        for j, c in enumerate(df.columns, start_col):
            cell = ws.cell(start_row, j, c)
            cell.fill = self.hdr_fill
            cell.font = self.hdr_font
            cell.alignment = self.center
            cell.border = self.border
        for i, row in enumerate(df.itertuples(index=False), start_row + 1):
            for j, val in enumerate(row, start_col):
                cell = ws.cell(i, j, None if pd.isna(val) else val)
                cell.font = self.body_font
                cell.alignment = self.right if isinstance(val, (int, float, np.integer, np.floating)) else self.left
                cell.border = self.border
        end_row = start_row + len(df)
        end_col = start_col + len(df.columns) - 1
        ref = f'{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}'
        tab = Table(displayName=table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)
        for c in range(start_col, end_col + 1):
            width = max(14, min(28, max(len(str(ws.cell(r, c).value)) if ws.cell(r, c).value is not None else 0 for r in range(start_row, min(end_row, start_row + 50) + 1)) + 2))
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.freeze_panes = f'{get_column_letter(start_col)}{start_row + 1}'
        return end_row, end_col

    def _add_actual_pred_chart(self, ws, start_row, actual_col, predicted_col, end_row):
        chart = ScatterChart()
        chart.title = 'Actual vs Predicted'
        chart.x_axis.title = 'Actual'
        chart.y_axis.title = 'Predicted'
        x = Reference(ws, min_col=actual_col, min_row=start_row + 1, max_row=end_row)
        y = Reference(ws, min_col=predicted_col, min_row=start_row + 1, max_row=end_row)
        chart.series.append(Series(y, x, title='Predicted'))
        chart.width = 14
        chart.height = 8
        ws.add_chart(chart, 'K2')

    def _add_output_chart(self, ws, start_row, end_row):
        chart = BarChart()
        chart.title = 'R2 by Split'
        chart.y_axis.title = 'R2'
        chart.x_axis.title = 'Split'
        chart.add_data(Reference(ws, min_col=5, min_row=start_row, max_row=end_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=2, min_row=start_row + 1, max_row=end_row))
        chart.width = 14
        chart.height = 8
        ws.add_chart(chart, 'H2')

    def write_excel(self):
        wb = Workbook()
        ws_overview = wb.active
        ws_overview.title = 'Overview'
        ws_eda = wb.create_sheet('EDA')
        ws_avp = wb.create_sheet('Actuals_vs_Predicted')
        ws_output = wb.create_sheet('Output')

        for ws in wb.worksheets:
            ws.column_dimensions['A'].width = 3

        ws_overview['B2'] = 'ML Model Report'
        ws_overview['B2'].font = self.title_font
        self._write_df(ws_overview, self.summary_df, 4, 2, 'OverviewTbl')

        r, _ = self._write_df(ws_eda, self.eda_missing_df, 2, 2, 'EDAMissingTbl')
        self._write_df(ws_eda, self.eda_stats_df.fillna(''), r + 3, 2, 'EDAStatsTbl')

        ar, _ = self._write_df(ws_avp, self.actual_pred_df.round(6), 2, 2, 'ActualPredTbl')
        actual_col = list(self.actual_pred_df.columns).index('Actual') + 2
        predicted_col = list(self.actual_pred_df.columns).index('Predicted') + 2
        self._add_actual_pred_chart(ws_avp, 2, actual_col, predicted_col, ar)

        orow, _ = self._write_df(ws_output, self.output_metrics_df.round(6), 2, 2, 'OutputTbl')
        self._add_output_chart(ws_output, 2, orow)

        wb.save(self.output_excel)
        return self.output_excel

    def run(self):
        self.load_data()
        self.build_model()
        self.prepare_sheets_data()
        self.save_all_plots()
        return self.write_excel()


if __name__ == '__main__':
    writer = ExcelWriterML(
        file_path='paddydataset.csv',
        output_excel='ml_model_report.xlsx',
        output_dir='plots',
        target='Nursery area (Cents)'
    )
    output_file = writer.run()
    print(f'Excel report created: {output_file}')
    print('Saved PNG plots:')
    for p in writer.plot_files:
        print(p)